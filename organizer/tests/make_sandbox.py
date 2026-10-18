"""Create a throwaway sandbox: destination folders with a few existing files, and an inbox to sort."""
import json
import shutil
import sys
from pathlib import Path

import docx
from openpyxl import Workbook

root = Path(sys.argv[1])
if root.exists():
    shutil.rmtree(root)
dest = {n: root / "Documents" / n for n in ("Finance", "School", "Research", "Code", "Career", "Personal")}
for p in dest.values():
    p.mkdir(parents=True)
inbox = root / "Inbox"
inbox.mkdir()

# Existing, already-organised files (these teach the similarity tier).
(dest["Finance"] / "2025-09 bank statement.txt").write_text("Account statement. Opening balance 1,240.00 USD. Deposits, withdrawals, card payments. Closing balance.")
(dest["School"] / "CSCI567 lecture 3 notes.md").write_text("# Lecture 3: linear regression\nGradient descent, loss functions, homework 2 due Friday. Midterm syllabus.")
(dest["Research"] / "experiment log.md").write_text("Experiment 12: NPU quantization ablation. Hypothesis, method, results table, perplexity, dataset.")
(dest["Career"] / "resume_2025.txt").write_text("Experience, Education, Skills. Software engineer intern. References available on request.")
(dest["Personal"] / "passport renewal.txt").write_text("Passport renewal appointment, travel dates, visa, apartment lease copy.")

# Inbox: a mix of easy (rules), medium (similarity) and ambiguous files.
(inbox / "Invoice_ACME_0042.txt").write_text("INVOICE #0042\nBill to: Soorya\nAmount due: $120.00\nPayment terms: net 30")
(inbox / "train_model.py").write_text("import torch\nfrom torch import nn\n\ndef train(model, data):\n    pass\n")
(inbox / "notes_week5.md").write_text("Lecture 5 notes: decision trees, entropy, homework 3 due next week. Quiz on Friday.")
(inbox / "ablation_results.csv").write_text("run,quantization,perplexity\n1,16a4w,12.3\n2,16a8w,11.1\n")
w = Workbook(); ws = w.active; ws.title = "Expenses"
ws.append(["Date", "Item", "Amount"]); ws.append(["2025-10-01", "Rent", 1500]); ws.append(["2025-10-03", "Groceries", 82])
w.save(inbox / "october.xlsx")
d = docx.Document()
d.add_paragraph("Dear Hiring Manager,")
d.add_paragraph("I am excited to apply for the Machine Learning Engineer position. My experience with on-device AI...")
d.add_paragraph("Sincerely, Soorya")
d.save(inbox / "letter.docx")
(inbox / "untitled.txt").write_text("Meeting at 3pm, bring the lease documents and the visa forms for the apartment.")
shutil.copy(dest["Finance"] / "2025-09 bank statement.txt", inbox / "statement copy.txt")  # duplicate

cfg = json.loads((Path(__file__).parent.parent / "config.example.json").read_text())
cfg["sources"] = [str(inbox)]
for n in cfg["destinations"]:
    cfg["destinations"][n]["path"] = str(dest[n])
cfg["state_dir"] = str(root / "state")
(root / "config.json").write_text(json.dumps(cfg, indent=1))
print(f"sandbox ready: {root} ({len(list(inbox.iterdir()))} inbox files)")
